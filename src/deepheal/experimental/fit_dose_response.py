#!/usr/bin/env python
from __future__ import annotations

from deepheal.paths import ARTICLE_ROOT


import argparse
import json
import math
import re
import string
import textwrap
from dataclasses import dataclass
from pathlib import Path

import matplotlib.pyplot as plt
import numpy as np
import pandas as pd
import seaborn as sns
from matplotlib.ticker import FixedLocator, FuncFormatter, NullFormatter
from openpyxl import load_workbook
from scipy.optimize import curve_fit
from scipy.stats import spearmanr


def find_project_root(start: Path) -> Path:
    return ARTICLE_ROOT


ROOT = ARTICLE_ROOT
RAW_XLSX = ROOT / "data" / "raw" / "cck_dose_response.xlsx"
TABLE_DIR = ROOT / "results" / "experimental" / "tables"
FIG_DIR = ROOT / "results" / "experimental" / "figures"
OUTPUT_ROOT = ROOT / "results" / "experimental"

RAW_LONG_PATH = TABLE_DIR / 'deepheal_cck8_ic50_raw_long.csv'
FIT_SUMMARY_PATH = TABLE_DIR / 'deepheal_cck8_ic50_fit_summary.csv'
UNDER10_SUMMARY_PATH = TABLE_DIR / 'deepheal_cck8_ic50_ic50_lt10uM_summary.csv'
SOURCE_BLOCK_SELECTION_PATH = TABLE_DIR / 'deepheal_cck8_ic50_source_block_selection.csv'
MANIFEST_PATH = TABLE_DIR / 'deepheal_cck8_ic50_manifest.json'

ALL18_STEM = FIG_DIR / 'deepheal_cck8_ic50_all18'
UNDER10_STEM = FIG_DIR / 'deepheal_cck8_ic50_ic50_lt10uM'

DPI = 300
MICRO = "\N{MICRO SIGN}"
RNG_SEED = 20260518
BOOTSTRAPS = 100
MIN_BOOTSTRAP_SUCCESSES = 60
MIN_R2 = 0.70
MIN_OBSERVED_SPAN = 15.0
MIN_FITTED_SPAN = 15.0
MIN_SPEARMAN_R = 0.35
MAX_CI_FOLD_WIDTH = 100.0


def configure_output_root(output_root: Path) -> None:
    global OUTPUT_ROOT, TABLE_DIR, FIG_DIR
    global RAW_LONG_PATH, FIT_SUMMARY_PATH, UNDER10_SUMMARY_PATH, SOURCE_BLOCK_SELECTION_PATH, MANIFEST_PATH
    global ALL18_STEM, UNDER10_STEM

    OUTPUT_ROOT = output_root.resolve()
    TABLE_DIR = OUTPUT_ROOT / "tables"
    FIG_DIR = OUTPUT_ROOT / "figures"
    RAW_LONG_PATH = TABLE_DIR / 'deepheal_cck8_ic50_raw_long.csv'
    FIT_SUMMARY_PATH = TABLE_DIR / 'deepheal_cck8_ic50_fit_summary.csv'
    UNDER10_SUMMARY_PATH = TABLE_DIR / 'deepheal_cck8_ic50_ic50_lt10uM_summary.csv'
    SOURCE_BLOCK_SELECTION_PATH = TABLE_DIR / 'deepheal_cck8_ic50_source_block_selection.csv'
    MANIFEST_PATH = TABLE_DIR / 'deepheal_cck8_ic50_manifest.json'
    ALL18_STEM = FIG_DIR / 'deepheal_cck8_ic50_all18'
    UNDER10_STEM = FIG_DIR / 'deepheal_cck8_ic50_ic50_lt10uM'


def relative_to_base(path: Path, base: Path) -> str:
    try:
        return str(path.resolve().relative_to(base.resolve()))
    except ValueError:
        return str(path)


def output_relative(path: Path) -> str:
    return relative_to_base(path, OUTPUT_ROOT)


@dataclass(frozen=True)
class CompoundBlock:
    sheet: str
    header_row: int
    concentration_col: int
    response_start_col: int
    response_end_col: int
    name: str
    unit: str


def normalize_text(value: object) -> str:
    if value is None:
        return ""
    return re.sub(r"\s+", " ", str(value).replace("\xa0", " ").strip())


def clean_compound_label(value: object) -> str:
    text = normalize_text(value)
    return re.sub(r"^\d+\.\s*", "", text).strip()


def display_unit(unit: str) -> str:
    return f"{MICRO}M" if unit == "uM" else unit


def parse_unit(value: object) -> str:
    text = normalize_text(value).lower()
    if "nm" in text:
        return "nM"
    if "um" in text or "micro" in text:
        return "uM"
    if MICRO.lower() in text or "\N{GREEK SMALL LETTER MU}m" in text:
        return "uM"
    return ""


def is_concentration_header(value: object) -> bool:
    return "concentration" in normalize_text(value).lower()


def is_numeric(value: object) -> bool:
    return isinstance(value, (int, float)) and math.isfinite(float(value))


def is_compound_header(value: object) -> bool:
    text = normalize_text(value)
    if not text or is_concentration_header(text) or text == "无效":
        return False
    if parse_unit(text):
        return False
    if re.fullmatch(r"[-+]?\d+(?:\.\d+)?\s*(?:uM|um|nM|nm)?", text, flags=re.IGNORECASE):
        return False
    return bool(re.match(r"^\d+\.\s*\S+", text))


def compound_order(name: str) -> tuple[int, str]:
    match = re.match(r"^(\d+)\.", normalize_text(name))
    return (int(match.group(1)) if match else 999, normalize_text(name))


def aligned_concentration_count(
    ws,
    header_row: int,
    concentration_col: int,
    response_start_col: int,
    response_end_col: int,
) -> int:
    aligned = 0
    seen_response_rows = False
    for row in range(header_row + 1, ws.max_row + 1):
        response_present = any(
            is_numeric(ws.cell(row, col).value) for col in range(response_start_col, response_end_col + 1)
        )
        if response_present:
            seen_response_rows = True
            if is_numeric(ws.cell(row, concentration_col).value):
                aligned += 1
        elif seen_response_rows:
            break
    return aligned


def find_unit_for_col(ws, header_row: int, concentration_col: int) -> str:
    direct = parse_unit(ws.cell(header_row, concentration_col).value)
    if direct:
        return direct
    for col in range(concentration_col - 1, 0, -1):
        if is_concentration_header(ws.cell(header_row, col).value):
            unit = parse_unit(ws.cell(header_row, col).value)
            if unit:
                return unit
    for col in range(concentration_col + 1, ws.max_column + 1):
        if is_concentration_header(ws.cell(header_row, col).value):
            unit = parse_unit(ws.cell(header_row, col).value)
            if unit:
                return unit
    return "uM"


def detect_compound_blocks(xlsx_path: Path) -> list[CompoundBlock]:
    wb = load_workbook(xlsx_path, data_only=True, read_only=False)
    blocks: list[CompoundBlock] = []

    for ws in wb.worksheets:
        merged_ranges = sorted(ws.merged_cells.ranges, key=lambda r: (r.min_row, r.min_col))
        for merged in merged_ranges:
            header_value = ws.cell(merged.min_row, merged.min_col).value
            if not is_compound_header(header_value):
                continue

            header_row = merged.min_row
            response_start = merged.min_col
            response_end = merged.max_col
            candidate_conc = response_start - 1

            if (
                candidate_conc >= 1
                and aligned_concentration_count(ws, header_row, candidate_conc, response_start, response_end) >= 2
            ):
                concentration_col = candidate_conc
            else:
                concentration_col = 0
                for col in range(response_start - 1, 0, -1):
                    if is_concentration_header(ws.cell(header_row, col).value):
                        concentration_col = col
                        break
                if concentration_col == 0:
                    raise ValueError(f"Could not find concentration column for {header_value!r} in {ws.title}")

            unit = find_unit_for_col(ws, header_row, concentration_col)
            blocks.append(
                CompoundBlock(
                    sheet=ws.title,
                    header_row=header_row,
                    concentration_col=concentration_col,
                    response_start_col=response_start,
                    response_end_col=response_end,
                    name=normalize_text(header_value),
                    unit=unit,
                )
            )

    return sorted(blocks, key=lambda block: compound_order(block.name))


def extract_existing_note(ws, block: CompoundBlock, last_data_row: int) -> str:
    notes: list[str] = []
    for row in range(last_data_row + 1, min(ws.max_row, last_data_row + 4) + 1):
        for col in range(block.response_start_col, block.response_end_col + 1):
            value = ws.cell(row, col).value
            text = normalize_text(value)
            if not text:
                continue
            if "ic50" in text.lower() or re.search(r"\d", text) and re.search(r"(?:uM|um|nM|nm|µm)", text, re.I):
                notes.append(text)
    return "; ".join(dict.fromkeys(notes))


def parse_excel_raw(xlsx_path: Path) -> pd.DataFrame:
    wb = load_workbook(xlsx_path, data_only=True, read_only=False)
    blocks = detect_compound_blocks(xlsx_path)
    rows: list[dict[str, object]] = []

    for block in blocks:
        ws = wb[block.sheet]
        last_data_row = block.header_row
        block_rows: list[dict[str, object]] = []

        for row_idx in range(block.header_row + 1, ws.max_row + 1):
            concentration = ws.cell(row_idx, block.concentration_col).value
            response_values = [
                ws.cell(row_idx, col).value for col in range(block.response_start_col, block.response_end_col + 1)
            ]
            numeric_responses = [value for value in response_values if is_numeric(value)]

            if is_numeric(concentration) and numeric_responses:
                last_data_row = row_idx
                for replicate_index, response in enumerate(response_values, start=1):
                    if not is_numeric(response):
                        continue
                    block_rows.append(
                        {
                            "compound": block.name,
                            "compound_label": clean_compound_label(block.name),
                            "compound_order": compound_order(block.name)[0],
                            "source_sheet": block.sheet,
                            "header_row": block.header_row,
                            "concentration_col": block.concentration_col,
                            "response_col": block.response_start_col + replicate_index - 1,
                            "unit": block.unit,
                            "concentration_native": float(concentration),
                            "concentration_uM": float(concentration) / 1000.0 if block.unit == "nM" else float(concentration),
                            "replicate": replicate_index,
                            "response": float(response),
                        }
                    )
            elif block_rows:
                break

        if not block_rows:
            raise ValueError(f"No numeric data parsed for {block.name}")

        existing_note = extract_existing_note(ws, block, last_data_row)
        for row in block_rows:
            row["existing_ic50_note"] = existing_note
        rows.extend(block_rows)

    raw = pd.DataFrame(rows)
    raw = raw.sort_values(["compound_order", "concentration_native", "replicate"]).reset_index(drop=True)
    return raw


def four_parameter_logistic(concentration: np.ndarray, bottom: float, span: float, log_ic50: float, hill: float) -> np.ndarray:
    concentration = np.asarray(concentration, dtype=float)
    exponent = np.clip((log_ic50 - np.log10(concentration)) * hill, -100.0, 100.0)
    return bottom + span / (1.0 + np.power(10.0, exponent))


def fit_4pl(
    concentration: np.ndarray,
    response: np.ndarray,
    start_params: np.ndarray | None = None,
    exhaustive: bool = True,
) -> tuple[np.ndarray | None, float]:
    concentration = np.asarray(concentration, dtype=float)
    response = np.asarray(response, dtype=float)
    mask = np.isfinite(concentration) & np.isfinite(response) & (concentration > 0)
    concentration = concentration[mask]
    response = response[mask]
    if concentration.size < 6 or np.unique(concentration).size < 4:
        return None, math.inf

    log_min = float(np.log10(np.min(concentration)))
    log_max = float(np.log10(np.max(concentration)))
    response_min = float(np.min(response))
    response_max = float(np.max(response))
    observed_span = max(response_max - response_min, 1.0)
    bounds = (
        [response_min - 60.0, 1.0e-6, log_min - 2.0, 0.05],
        [response_max + 80.0, 250.0, log_max + 2.0, 8.0],
    )

    starts: list[list[float]] = []
    if start_params is not None:
        base = np.asarray(start_params, dtype=float)
        starts.append(base.tolist())
    elif exhaustive:
        bottom_starts = [np.percentile(response, q) for q in (0, 10, 25)]
        span_starts = [observed_span, max(observed_span * 1.5, 20.0), 100.0]
        log_starts = np.linspace(log_min, log_max, 5)
        hill_starts = [0.25, 0.5, 1.0, 2.0, 4.0]
        for bottom in bottom_starts:
            for span in span_starts:
                for log_ic50 in log_starts:
                    for hill in hill_starts:
                        starts.append([float(bottom), float(span), float(log_ic50), float(hill)])
    else:
        bottom_starts = [np.percentile(response, q) for q in (0, 10)]
        span_starts = [observed_span, max(observed_span * 1.3, 20.0)]
        log_starts = [float(np.median([log_min, log_max])), log_min, log_max]
        hill_starts = [0.5, 1.0, 2.0]
        for bottom in bottom_starts:
            for span in span_starts:
                for log_ic50 in log_starts:
                    for hill in hill_starts:
                        starts.append([float(bottom), float(span), float(log_ic50), float(hill)])

    best_params: np.ndarray | None = None
    best_rss = math.inf
    for start in starts:
        clipped_start = np.clip(np.asarray(start, dtype=float), bounds[0], bounds[1])
        try:
            params, _ = curve_fit(
                four_parameter_logistic,
                concentration,
                response,
                p0=clipped_start,
                bounds=bounds,
                method="trf",
                maxfev=8000 if start_params is not None else 40000,
            )
        except (RuntimeError, ValueError, FloatingPointError):
            continue
        fitted = four_parameter_logistic(concentration, *params)
        rss = float(np.sum((response - fitted) ** 2))
        if rss < best_rss:
            best_rss = rss
            best_params = params

    return best_params, best_rss


def bootstrap_ic50(
    compound_df: pd.DataFrame,
    rng: np.random.Generator,
    start_params: np.ndarray,
    n_bootstrap: int = BOOTSTRAPS,
) -> tuple[float, float, int, float]:
    grouped = [
        (float(conc), sub["response"].to_numpy(dtype=float))
        for conc, sub in compound_df.groupby("concentration_native", sort=True)
    ]
    boot_ic50: list[float] = []

    for _ in range(n_bootstrap):
        concentrations: list[float] = []
        responses: list[float] = []
        for concentration, values in grouped:
            resampled = rng.choice(values, size=len(values), replace=True)
            concentrations.extend([concentration] * len(resampled))
            responses.extend(resampled.tolist())

        params, _ = fit_4pl(
            np.asarray(concentrations, dtype=float),
            np.asarray(responses, dtype=float),
            start_params=start_params,
            exhaustive=False,
        )
        if params is None:
            continue
        ic50 = float(10.0 ** params[2])
        if math.isfinite(ic50) and ic50 > 0:
            boot_ic50.append(ic50)

    if len(boot_ic50) < MIN_BOOTSTRAP_SUCCESSES:
        return math.nan, math.nan, len(boot_ic50), math.nan

    low, high = np.percentile(np.asarray(boot_ic50, dtype=float), [2.5, 97.5])
    fold_width = float(high / low) if low > 0 else math.inf
    return float(low), float(high), len(boot_ic50), fold_width


def summarize_fit(compound: str, compound_df: pd.DataFrame, rng: np.random.Generator) -> dict[str, object]:
    unit = str(compound_df["unit"].iloc[0])
    concentrations = compound_df["concentration_native"].to_numpy(dtype=float)
    responses = compound_df["response"].to_numpy(dtype=float)
    x_unique = np.sort(compound_df["concentration_native"].unique().astype(float))
    mean_by_conc = compound_df.groupby("concentration_native", sort=True)["response"].mean()

    params, rss = fit_4pl(concentrations, responses)
    observed_span = float(mean_by_conc.max() - mean_by_conc.min())
    spearman_value = float(spearmanr(np.log10(mean_by_conc.index.to_numpy(dtype=float)), mean_by_conc.to_numpy()).statistic)

    summary: dict[str, object] = {
        "compound": compound,
        "compound_label": clean_compound_label(compound),
        "compound_order": int(compound_df["compound_order"].iloc[0]),
        "unit": unit,
        "n_concentrations": int(x_unique.size),
        "n_observations": int(len(compound_df)),
        "replicate_count_min": int(compound_df.groupby("concentration_native").size().min()),
        "replicate_count_max": int(compound_df.groupby("concentration_native").size().max()),
        "concentration_min_native": float(np.min(x_unique)),
        "concentration_max_native": float(np.max(x_unique)),
        "observed_response_min": float(mean_by_conc.min()),
        "observed_response_max": float(mean_by_conc.max()),
        "observed_span": observed_span,
        "spearman_r_log_conc_response": spearman_value,
        "existing_ic50_note": str(compound_df["existing_ic50_note"].iloc[0]),
    }

    reasons: list[str] = []
    if params is None:
        reasons.append("fit failed")
        summary.update(
            {
                "fitted_bottom": math.nan,
                "fitted_top": math.nan,
                "fitted_span": math.nan,
                "hill_slope": math.nan,
                "log_ic50": math.nan,
                "ic50_native": math.nan,
                "ic50_ci_low_native": math.nan,
                "ic50_ci_high_native": math.nan,
                "ic50_uM": math.nan,
                "ic50_ci_low_uM": math.nan,
                "ic50_ci_high_uM": math.nan,
                "r_squared": math.nan,
                "rmse": math.nan,
                "bootstrap_successes": 0,
                "bootstrap_ci_fold_width": math.nan,
            }
        )
    else:
        fitted = four_parameter_logistic(concentrations, *params)
        ss_total = float(np.sum((responses - np.mean(responses)) ** 2))
        r_squared = 1.0 - rss / ss_total if ss_total > 0 else math.nan
        rmse = math.sqrt(rss / max(len(responses) - len(params), 1))
        bottom, span, log_ic50, hill = [float(value) for value in params]
        fitted_top = bottom + span
        ic50_native = float(10.0 ** log_ic50)
        unit_factor = 0.001 if unit == "nM" else 1.0

        if not (np.min(x_unique) <= ic50_native <= np.max(x_unique)):
            reasons.append("IC50 outside tested range")
        if not math.isfinite(r_squared) or r_squared < MIN_R2:
            reasons.append(f"low R2 ({r_squared:.3f})")
        if observed_span < MIN_OBSERVED_SPAN:
            reasons.append(f"small observed span ({observed_span:.1f})")
        if span < MIN_FITTED_SPAN:
            reasons.append(f"small fitted span ({span:.1f})")
        if not math.isfinite(spearman_value) or spearman_value < MIN_SPEARMAN_R:
            reasons.append(f"weak monotonicity (Spearman r={spearman_value:.3f})")

        ci_low = math.nan
        ci_high = math.nan
        boot_successes = 0
        ci_fold_width = math.nan
        if not reasons:
            ci_low, ci_high, boot_successes, ci_fold_width = bootstrap_ic50(compound_df, rng, params)
            if (
                not math.isfinite(ci_low)
                or not math.isfinite(ci_high)
                or ci_low <= 0
                or ci_high <= ci_low
                or boot_successes < MIN_BOOTSTRAP_SUCCESSES
                or not math.isfinite(ci_fold_width)
                or ci_fold_width > MAX_CI_FOLD_WIDTH
            ):
                reasons.append("unstable bootstrap CI")

        summary.update(
            {
                "fitted_bottom": bottom,
                "fitted_top": fitted_top,
                "fitted_span": span,
                "hill_slope": hill,
                "log_ic50": log_ic50,
                "ic50_native": ic50_native,
                "ic50_ci_low_native": ci_low,
                "ic50_ci_high_native": ci_high,
                "ic50_uM": ic50_native * unit_factor,
                "ic50_ci_low_uM": ci_low * unit_factor if math.isfinite(ci_low) else math.nan,
                "ic50_ci_high_uM": ci_high * unit_factor if math.isfinite(ci_high) else math.nan,
                "r_squared": r_squared,
                "rmse": rmse,
                "bootstrap_successes": int(boot_successes),
                "bootstrap_ci_fold_width": ci_fold_width,
            }
        )

    summary["fit_status"] = "reliable" if not reasons else "n.d."
    summary["nd_reason"] = "; ".join(reasons)
    return summary


def build_fit_summary(raw: pd.DataFrame) -> pd.DataFrame:
    rng = np.random.default_rng(RNG_SEED)
    summaries = [
        summarize_fit(compound, compound_df.copy(), rng)
        for compound, compound_df in raw.groupby("compound", sort=False)
    ]
    return pd.DataFrame(summaries).sort_values("compound_order").reset_index(drop=True)


def compact_tick(value: float, _pos: int | None = None) -> str:
    if value <= 0 or not math.isfinite(value):
        return ""
    if value >= 100:
        return f"{value:.0f}"
    if value >= 10:
        return f"{value:.0f}" if abs(value - round(value)) < 1e-8 else f"{value:.1f}"
    if value >= 1:
        return f"{value:.1f}".rstrip("0").rstrip(".")
    return f"{value:.3g}"


def format_ic50(row: pd.Series) -> str:
    unit = display_unit(str(row["unit"]))
    if row["fit_status"] != "reliable":
        return "IC50 n.d."
    low = float(row["ic50_ci_low_native"])
    high = float(row["ic50_ci_high_native"])
    ic50 = float(row["ic50_native"])
    if math.isfinite(low) and math.isfinite(high):
        return f"IC50 {compact_tick(ic50)} {unit}\n95% CI {compact_tick(low)}-{compact_tick(high)}"
    return f"IC50 {compact_tick(ic50)} {unit}"


def source_block_columns(raw: pd.DataFrame) -> pd.DataFrame:
    raw = raw.copy()
    block_bounds = raw.groupby(["compound", "source_sheet", "header_row", "concentration_col"])["response_col"].agg(
        response_start_col="min",
        response_end_col="max",
    )
    raw = raw.merge(
        block_bounds.reset_index(),
        on=["compound", "source_sheet", "header_row", "concentration_col"],
        how="left",
    )
    raw["source_block_key"] = (
        raw["source_sheet"].astype(str)
        + "|r"
        + raw["header_row"].astype(int).astype(str)
        + "|c"
        + raw["concentration_col"].astype(int).astype(str)
        + "|resp"
        + raw["response_start_col"].astype(int).astype(str)
        + "-"
        + raw["response_end_col"].astype(int).astype(str)
    )
    return raw


def build_source_block_selection(raw_all: pd.DataFrame) -> tuple[pd.DataFrame, pd.DataFrame]:
    raw_all = source_block_columns(raw_all)
    block_rows: list[dict[str, object]] = []
    for compound, compound_blocks in raw_all.groupby("compound", sort=False):
        block_summary = (
            compound_blocks.groupby(
                [
                    "source_block_key",
                    "compound",
                    "compound_label",
                    "compound_order",
                    "unit",
                    "source_sheet",
                    "header_row",
                    "concentration_col",
                    "response_start_col",
                    "response_end_col",
                ],
                sort=True,
            )
            .agg(
                n_concentrations=("concentration_native", "nunique"),
                n_observations=("response", "size"),
                concentration_min_native=("concentration_native", "min"),
                concentration_max_native=("concentration_native", "max"),
            )
            .reset_index()
            .sort_values(["source_sheet", "header_row", "concentration_col", "response_start_col"])
            .reset_index(drop=True)
        )
        block_summary["run_index"] = np.arange(1, len(block_summary) + 1)
        selected_run_index = 2 if len(block_summary) >= 2 else 1
        for row in block_summary.to_dict("records"):
            row["selected"] = int(row["run_index"]) == selected_run_index
            row["selection_rule"] = (
                "selected run #2; user specified SB225002 #2 and duplicate May 18 blocks must not be combined"
                if selected_run_index == 2
                else "selected sole detected run"
            )
            block_rows.append(row)

    selection = pd.DataFrame(block_rows).sort_values(["compound_order", "run_index"]).reset_index(drop=True)
    selected_keys = set(selection.loc[selection["selected"], "source_block_key"])
    raw_selected = raw_all.loc[raw_all["source_block_key"].isin(selected_keys)].copy()
    raw_selected = raw_selected.sort_values(["compound_order", "concentration_native", "replicate"]).reset_index(drop=True)
    return raw_selected, selection


def set_publication_style() -> None:
    plt.rcParams.update(
        {
            "font.family": "sans-serif",
            "font.sans-serif": ["Arial", "Helvetica", "DejaVu Sans"],
            "font.size": 7,
            "axes.labelsize": 7,
            "axes.titlesize": 7.4,
            "xtick.labelsize": 5.8,
            "ytick.labelsize": 5.8,
            "axes.linewidth": 0.65,
            "xtick.major.width": 0.55,
            "ytick.major.width": 0.55,
            "pdf.fonttype": 42,
            "ps.fonttype": 42,
            "figure.dpi": DPI,
            "savefig.dpi": DPI,
        }
    )
    sns.set_theme(style="ticks", context="paper")


def choose_layout(n_panels: int) -> tuple[int, int, tuple[float, float]]:
    if n_panels <= 6:
        return 2, 3, (7.2, 3.9)
    return math.ceil(n_panels / 3), 3, (7.2, 11.0)


def draw_dose_response_figure(
    raw: pd.DataFrame,
    fit_summary: pd.DataFrame,
    output_stem: Path,
    *,
    panel_labels: bool = True,
) -> list[Path]:
    set_publication_style()
    line_color = "#0072B2"
    mean_color = "#D55E00"
    point_color = "#4D4D4D"
    band_color = "#56B4E9"
    rng = np.random.default_rng(RNG_SEED)

    nrows, ncols, figsize = choose_layout(len(fit_summary))
    fig, axes = plt.subplots(nrows, ncols, figsize=figsize, sharey=False, squeeze=False)
    axes_flat = axes.ravel()

    for idx, (_, fit_row) in enumerate(fit_summary.iterrows()):
        ax = axes_flat[idx]
        compound = str(fit_row["compound"])
        compound_label = str(fit_row["compound_label"])
        sub = raw.loc[raw["compound"] == compound].copy()
        unit = str(fit_row["unit"])
        grouped = sub.groupby("concentration_native", sort=True)["response"]
        means = grouped.mean()
        sems = grouped.sem().fillna(0.0)
        x_values = means.index.to_numpy(dtype=float)

        for _, point in sub.iterrows():
            jitter = float(np.exp(rng.normal(0.0, 0.025)))
            ax.scatter(
                float(point["concentration_native"]) * jitter,
                float(point["response"]),
                s=8,
                c=point_color,
                alpha=0.52,
                linewidths=0,
                zorder=2,
            )

        ax.errorbar(
            x_values,
            means.to_numpy(dtype=float),
            yerr=sems.to_numpy(dtype=float),
            fmt="o",
            ms=3.0,
            mfc=mean_color,
            mec="white",
            mew=0.35,
            ecolor=mean_color,
            elinewidth=0.75,
            capsize=1.8,
            zorder=4,
        )

        if fit_row["fit_status"] == "reliable":
            x_curve = np.logspace(np.log10(x_values.min()), np.log10(x_values.max()), 240)
            params = [
                float(fit_row["fitted_bottom"]),
                float(fit_row["fitted_span"]),
                float(fit_row["log_ic50"]),
                float(fit_row["hill_slope"]),
            ]
            y_curve = four_parameter_logistic(x_curve, *params)
            ax.plot(x_curve, y_curve, color=line_color, lw=1.25, zorder=3)

            low = float(fit_row["ic50_ci_low_native"])
            high = float(fit_row["ic50_ci_high_native"])
            if math.isfinite(low) and math.isfinite(high):
                lower_params = params.copy()
                upper_params = params.copy()
                lower_params[2] = math.log10(low)
                upper_params[2] = math.log10(high)
                y_low = four_parameter_logistic(x_curve, *lower_params)
                y_high = four_parameter_logistic(x_curve, *upper_params)
                band_low = np.minimum(y_low, y_high)
                band_high = np.maximum(y_low, y_high)
                ax.fill_between(x_curve, band_low, band_high, color=band_color, alpha=0.16, linewidth=0, zorder=1)

        ax.set_xscale("log")
        ax.set_xlim(x_values.min() / 1.25, x_values.max() * 1.25)
        y_min = min(float(sub["response"].min()), float(means.min())) - 8
        y_max = max(float(sub["response"].max()), float(means.max())) + 12
        ax.set_ylim(max(-25, y_min), min(125, max(60, y_max)))
        ax.xaxis.set_major_locator(FixedLocator([x_values[0], x_values[len(x_values) // 2], x_values[-1]]))
        ax.xaxis.set_major_formatter(FuncFormatter(compact_tick))
        ax.xaxis.set_minor_formatter(NullFormatter())
        ax.tick_params(axis="both", length=2.2, pad=1.5)
        ax.grid(axis="y", color="#E6E6E6", linewidth=0.45)
        ax.grid(axis="x", visible=False)
        ax.spines["top"].set_visible(False)
        ax.spines["right"].set_visible(False)

        ax.set_title(textwrap.fill(compound_label, width=18), loc="left", pad=2.0, fontweight="bold")
        ax.text(
            0.03,
            0.95,
            format_ic50(fit_row),
            transform=ax.transAxes,
            va="top",
            ha="left",
            fontsize=5.7,
            color="#1A1A1A",
            bbox={"boxstyle": "round,pad=0.18", "facecolor": "white", "edgecolor": "none", "alpha": 0.78},
        )
        if panel_labels:
            ax.text(
                -0.18,
                1.08,
                string.ascii_lowercase[idx],
                transform=ax.transAxes,
                fontsize=8,
                fontweight="bold",
                va="top",
            )
        ax.set_xlabel(f"Concentration ({display_unit(unit)})", labelpad=1.8)
        if idx % ncols == 0:
            ax.set_ylabel("Response (%)", labelpad=1.8)
        else:
            ax.set_ylabel("")

    for empty_ax in axes_flat[len(fit_summary) :]:
        empty_ax.axis("off")

    if len(fit_summary) <= 6:
        fig.subplots_adjust(left=0.075, right=0.985, top=0.94, bottom=0.12, wspace=0.36, hspace=0.58)
    else:
        fig.subplots_adjust(left=0.075, right=0.985, top=0.975, bottom=0.055, wspace=0.36, hspace=0.68)

    output_paths: list[Path] = []
    for suffix in [".png", ".pdf"]:
        path = output_stem.with_suffix(suffix)
        fig.savefig(path, dpi=DPI, bbox_inches="tight")
        output_paths.append(path)
    plt.close(fig)
    return output_paths


def validate_outputs(raw: pd.DataFrame, fit_summary: pd.DataFrame, under10: pd.DataFrame, output_paths: list[Path]) -> None:
    if raw["compound"].nunique() != 18:
        raise AssertionError(f"Expected 18 compounds, parsed {raw['compound'].nunique()}")
    if len(fit_summary) != 18:
        raise AssertionError(f"Expected 18 fit rows, found {len(fit_summary)}")
    if not (raw["concentration_native"] > 0).all():
        raise AssertionError("All parsed concentrations must be positive.")
    if len(under10) == 0:
        raise AssertionError("No reliable compounds with IC50 < 10 uM were found.")
    if not (under10["fit_status"].eq("reliable") & (under10["ic50_uM"] < 10)).all():
        raise AssertionError("Filtered under-10 table contains non-reliable or >=10 uM rows.")
    block_counts = raw.groupby("compound")["source_block_key"].nunique()
    if int(block_counts.min()) != 1 or int(block_counts.max()) != 1:
        raise AssertionError(f"Expected exactly one selected source block per compound, observed {block_counts.to_dict()}")
    sb225002 = raw.loc[raw["compound_label"] == "SB225002"]
    if sb225002.empty:
        raise AssertionError("SB225002 was not parsed.")
    sb_block = sb225002[["source_sheet", "header_row", "concentration_col", "response_start_col", "response_end_col"]].drop_duplicates()
    expected_sb_block = ("Sheet1", 47, 11, 12, 14)
    observed_sb_blocks = [tuple(row) for row in sb_block.to_numpy()]
    if observed_sb_blocks != [expected_sb_block]:
        raise AssertionError(f"SB225002 must use #2 block {expected_sb_block}, observed {observed_sb_blocks}")

    for path in [RAW_LONG_PATH, FIT_SUMMARY_PATH, UNDER10_SUMMARY_PATH, SOURCE_BLOCK_SELECTION_PATH, MANIFEST_PATH, *output_paths]:
        if not path.exists() or path.stat().st_size <= 0:
            raise AssertionError(f"Missing or empty output: {path}")


def run_analysis() -> dict[str, object]:
    TABLE_DIR.mkdir(parents=True, exist_ok=True)
    FIG_DIR.mkdir(parents=True, exist_ok=True)

    raw_all = parse_excel_raw(RAW_XLSX)
    raw, source_block_selection = build_source_block_selection(raw_all)
    fit_summary = build_fit_summary(raw)
    under10 = fit_summary.loc[(fit_summary["fit_status"] == "reliable") & (fit_summary["ic50_uM"] < 10)].copy()

    raw.to_csv(RAW_LONG_PATH, index=False, encoding="utf-8-sig")
    fit_summary.to_csv(FIT_SUMMARY_PATH, index=False, encoding="utf-8-sig")
    under10.to_csv(UNDER10_SUMMARY_PATH, index=False, encoding="utf-8-sig")
    source_block_selection.to_csv(SOURCE_BLOCK_SELECTION_PATH, index=False, encoding="utf-8-sig")

    output_paths = []
    output_paths.extend(draw_dose_response_figure(raw, fit_summary, ALL18_STEM, panel_labels=True))
    output_paths.extend(draw_dose_response_figure(raw, under10, UNDER10_STEM, panel_labels=False))

    manifest = {
        "source_workbook": relative_to_base(RAW_XLSX, ROOT),
        "source_code_reused": "self-contained clean-copy May 18 CCK-8 parser and 4PL fitter",
        "raw_long_csv": output_relative(RAW_LONG_PATH),
        "fit_summary_csv": output_relative(FIT_SUMMARY_PATH),
        "ic50_lt10uM_summary_csv": output_relative(UNDER10_SUMMARY_PATH),
        "source_block_selection_csv": output_relative(SOURCE_BLOCK_SELECTION_PATH),
        "figures": [output_relative(path) for path in output_paths],
        "dpi": DPI,
        "compound_count": int(raw["compound"].nunique()),
        "all_detected_source_block_count": int(source_block_selection["source_block_key"].nunique()),
        "selected_source_block_count": int(source_block_selection["selected"].sum()),
        "observation_count": int(len(raw)),
        "fit_status_counts": {str(key): int(value) for key, value in fit_summary["fit_status"].value_counts().items()},
        "ic50_lt10uM_compounds": under10["compound_label"].tolist(),
        "selection_policy": "One source block per compound; duplicate May 18 blocks use run #2, including SB225002 #2.",
    }
    MANIFEST_PATH.write_text(json.dumps(manifest, ensure_ascii=False, indent=2), encoding="utf-8")
    validate_outputs(raw, fit_summary, under10, output_paths)
    return manifest


def parse_args() -> argparse.Namespace:
    parser = argparse.ArgumentParser(description="Generate May 18 CCK-8 IC50 tables and figures.")
    parser.add_argument(
        "--output-root",
        type=Path,
        default=ROOT / "results" / "experimental",
        help="Directory for generated tables/figures. Defaults to results/experimental.",
    )
    return parser.parse_args()


def main() -> None:
    args = parse_args()
    configure_output_root(args.output_root)
    manifest = run_analysis()
    print("May 18 CCK8 IC50 analysis completed")
    print(f"Parsed compounds: {manifest['compound_count']}")
    print(f"Raw observations: {manifest['observation_count']}")
    print(f"Fit status counts: {manifest['fit_status_counts']}")
    print(f"IC50 < 10 {MICRO}M: {', '.join(manifest['ic50_lt10uM_compounds'])}")
    print("Outputs:")
    for path in [
        RAW_LONG_PATH,
        FIT_SUMMARY_PATH,
        UNDER10_SUMMARY_PATH,
        SOURCE_BLOCK_SELECTION_PATH,
        MANIFEST_PATH,
        *map(Path, manifest["figures"]),
    ]:
        print(f"  {path} ({path.stat().st_size} bytes)")


if __name__ == "__main__":
    main()
